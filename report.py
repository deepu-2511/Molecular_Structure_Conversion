from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out = Path(__file__).resolve().parent / "MolecularStructureAI_Detailed_Project_Report.docx"
doc = Document()

# Page setup
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)

for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
    styles[style_name].font.name = "Aptos"

# Helpers
def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(9)

def add_table(headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, True)
        shade_cell(hdr[i], "D9E2F3")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return table

def add_bullets(items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)

def add_code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.25)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    return p

# Title
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("MolecularStructureAI")
r.bold = True
r.font.size = Pt(24)
r.font.name = "Aptos Display"

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Detailed Technical Project Report")
r.bold = True
r.font.size = Pt(15)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Molecular structure generation, validation, visualization, analysis and export platform").italic = True

doc.add_paragraph()

# Executive summary
doc.add_heading("1. Executive Summary", level=1)
doc.add_paragraph(
    "MolecularStructureAI is a research-oriented molecular structure platform designed to accept "
    "chemical representations and convert them into validated molecular structures with 2D/3D "
    "visualization, molecular descriptors, identifiers and downloadable structure files. "
    "The implemented architecture separates deterministic cheminformatics operations from future "
    "machine-learning components. RDKit is the core chemistry engine, FastAPI provides the backend API, "
    "and a React/Vite frontend provides the user interface."
)
doc.add_paragraph(
    "At the current stage, the SMILES and IUPAC workflows are operational. SMILES input is parsed and "
    "validated using RDKit. IUPAC input is resolved through PubChem PUG REST and then passed through the "
    "same RDKit molecular pipeline. The system generates 2D PNG/SVG output, ETKDGv3-based 3D conformers, "
    "MMFF/UFF optimization, SDF/MOL/PDB/XYZ files, molecular properties, identifiers and job metadata."
)

# Current status
doc.add_heading("2. Current Development Status", level=1)
add_table(
    ["Module", "Status", "Current implementation"],
    [
        ["Project structure", "Complete", "Backend, chemistry modules, frontend, ML and export directories established."],
        ["SMILES parser", "Complete", "RDKit MolFromSmiles + sanitization."],
        ["Molecular descriptors", "Complete", "Formula, MW, exact mass, LogP, TPSA, HBD/HBA, rotatable bonds, rings, charge."],
        ["Chemical identifiers", "Complete", "Canonical/isomeric SMILES, InChI, InChIKey."],
        ["2D rendering", "Complete", "RDKit PNG and SVG rendering."],
        ["3D generation", "Complete", "ETKDGv3 embedding + MMFF, UFF fallback."],
        ["3D structure export", "Complete", "SDF, MOL, PDB, XYZ."],
        ["FastAPI backend", "Complete", "Generation, retrieval, download and JSON endpoints."],
        ["IUPAC resolution", "Working", "PubChem PUG REST resolver; robust synonym/multiple-match handling remains to be hardened."],
        ["React frontend", "Working", "Old scientific workstation UI retained and connected to API."],
        ["Image → structure ML", "Planned", "Dedicated ML module not yet integrated."],
        ["Interactive 3D viewer", "Planned", "Future frontend module."],
        ["Extended exports", "Planned", "CML, PDF, TIFF, WEBP, CSV, ZIP."],
        ["Validation/confidence", "Planned", "Chemical validation and confidence/explainability layer."],
        ["History/library", "Planned", "Persistent jobs and structure library."],
        ["Production hardening", "Planned", "Authentication, rate limits, job queue, caching, logging, deployment."],
    ],
    [1.8, 1.0, 4.7]
)

# Objectives
doc.add_heading("3. System Objectives", level=1)
add_bullets([
    "Accept SMILES, chemical/IUPAC names and, in the future, chemical structure images.",
    "Convert input into a validated RDKit molecular graph.",
    "Generate chemically meaningful 2D and computational 3D representations.",
    "Calculate standard molecular descriptors and machine-readable chemical identifiers.",
    "Provide downloadable industry/research-compatible molecular formats.",
    "Keep deterministic cheminformatics separate from machine-learning inference.",
    "Provide an API-first architecture so the frontend is not coupled to the chemistry implementation.",
    "Create a foundation that can later support image recognition, confidence scoring, provenance and structure libraries."
])

# Architecture
doc.add_heading("4. Overall System Architecture", level=1)
add_code(
"""                         MolecularStructureAI
                                │
          ┌─────────────────────┴─────────────────────┐
          │                                           │
      React/Vite                                  FastAPI
      Frontend                                    Backend API
          │                                           │
          │                                  ┌────────┴────────┐
          │                                  │ MoleculeEngine  │
          │                                  └────────┬────────┘
          │                                           │
          │                  ┌────────────────────────┼───────────────────────┐
          │                  │                        │                       │
          │              RDKit Core              2D Renderer             3D Engine
          │                  │                        │                       │
          │             Validation +             PNG/SVG                 ETKDGv3
          │             descriptors                                      MMFF/UFF
          │                  │                        │                       │
          │                  └────────────────────────┼───────────────────────┘
          │                                           │
          │                                  Export / Metadata
          │                                           │
          └───────────────────────────────────────────┘
                               │
                     Future Image → Structure ML
"""
)
doc.add_paragraph(
    "The central design principle is that every supported input method should eventually converge on "
    "one validated molecular graph. After that point, 2D rendering, 3D generation, descriptors and exports "
    "should use the same chemistry engine rather than duplicating logic."
)

# Input workflows
doc.add_heading("5. Input Workflows", level=1)
doc.add_heading("5.1 SMILES Workflow", level=2)
add_code(
"""User SMILES
    ↓
MoleculeEngine.from_smiles()
    ↓
Chem.MolFromSmiles()
    ↓
Chem.SanitizeMol()
    ↓
Validated RDKit Mol
    ↓
Identifiers + descriptors
    ↓
Optional complete generation
    ├── 2D PNG/SVG
    ├── 3D conformer
    ├── SDF/MOL/PDB/XYZ
    └── metadata.json
"""
)

doc.add_heading("5.2 IUPAC Workflow", level=2)
add_code(
"""User chemical name
    ↓
IUPACResolver
    ↓
PubChem PUG REST
    ↓
CID + returned molecular representation
    ↓
MoleculeEngine.from_smiles()
    ↓
RDKit validation
    ↓
Same complete-generation pipeline as SMILES
"""
)
doc.add_paragraph(
    "Important design point: the IUPAC resolver does not itself generate the final molecular structure. "
    "It resolves a name to a molecular representation, after which RDKit becomes the validation and structure "
    "generation authority inside the application."
)

doc.add_heading("5.3 Future Image Workflow", level=2)
add_code(
"""Structure image
    ↓
Image preprocessing
    ↓
Image-to-structure ML model
    ↓
Predicted molecular representation
    ↓
RDKit parsing + sanitization
    ↓
Chemical validation
    ↓
Confidence / uncertainty checks
    ↓
2D + 3D + descriptors + exports
"""
)

# Backend modules
doc.add_heading("6. Backend Modules — Inch-by-Inch", level=1)

doc.add_heading("6.1 backend/main.py", level=2)
doc.add_paragraph(
    "Acts as the FastAPI application layer. It initializes the application, configures CORS, exposes "
    "HTTP endpoints, creates generation jobs, stores output files, and serves generated files."
)
add_bullets([
    "Creates the FastAPI application with title, description and version.",
    "Allows the React development server at localhost:5173 to communicate with the API.",
    "Creates the exports directory if it does not exist.",
    "Mounts /exports as a static directory.",
    "Defines request models for SMILES and IUPAC name inputs.",
    "Uses one shared generate_complete_structure() function so SMILES and IUPAC do not duplicate generation logic.",
    "Creates a UUID job directory for every complete-generation request.",
    "Stores metadata.json as the machine-readable record of a generated structure.",
    "Provides job retrieval and file download endpoints."
])

doc.add_heading("6.2 backend/chemistry/core.py", level=2)
doc.add_paragraph(
    "This is the deterministic chemistry core and the main abstraction around RDKit."
)
add_bullets([
    "from_smiles(): parses a SMILES string and sanitizes the resulting RDKit molecule.",
    "get_properties(): calculates formula, molecular weight, exact mass, LogP, TPSA, HBD, HBA, rotatable bonds, heavy atoms, rings, aromatic rings and formal charge.",
    "get_identifiers(): generates canonical SMILES, isomeric SMILES, InChI and InChIKey.",
    "analyze_smiles(): provides a compact analysis response combining identifiers and properties."
])

doc.add_heading("6.3 backend/chemistry/structure_2d.py", level=2)
doc.add_paragraph(
    "Responsible for deterministic 2D depiction using RDKit."
)
add_bullets([
    "Copies the molecule before rendering.",
    "Computes 2D coordinates using RDKit.",
    "Renders PNG output through Draw.MolToImage().",
    "Renders vector SVG output through MolDraw2DSVG.",
    "Creates output directories automatically."
])

doc.add_heading("6.4 backend/chemistry/structure_3d.py", level=2)
doc.add_paragraph(
    "Generates a computational 3D conformer and exports standard 3D molecular files."
)
add_bullets([
    "Adds explicit hydrogens before 3D embedding.",
    "Uses ETKDGv3 with a fixed random seed for reproducibility.",
    "Attempts MMFF optimization first.",
    "Falls back to UFF if MMFF cannot optimize.",
    "Reports the optimization method.",
    "Calculates MMFF energy where supported.",
    "Exports SDF, MOL, PDB and XYZ.",
    "The generated 3D geometry is computational and should not be represented as experimental structure data."
])

doc.add_heading("6.5 backend/chemistry/iupac.py", level=2)
doc.add_paragraph(
    "Resolves chemical names through PubChem PUG REST and returns molecular identifiers/SMILES. "
    "The current implementation is functional for tested names, but the resolver should be upgraded "
    "before production deployment to handle naming variants, synonyms and ambiguous matches."
)
add_bullets([
    "Validates that the name is not empty.",
    "URL-encodes the name.",
    "Queries PubChem PUG REST name/property resolution.",
    "Extracts CID, IUPAC name, SMILES, InChI and InChIKey where returned.",
    "Converts the resolved SMILES into an RDKit molecule.",
    "Raises a controlled error when resolution fails."
])

# Resolver upgrade
doc.add_heading("7. IUPAC Resolver — Required Production Upgrade", level=1)
doc.add_paragraph(
    "The current exact-name resolver is sufficient for basic operation but is not enough for a research-grade "
    "chemical name interface. Researchers may copy names from PubChem, ChEMBL, DrugBank, papers, patents or "
    "vendor databases, where nomenclature style, synonyms and formatting can differ while referring to the same structure."
)
add_code(
"""Recommended resolution strategy:

Input name
   ↓
Normalize whitespace / Unicode / punctuation
   ↓
1. Exact PubChem name lookup
   ↓
2. PubChem synonym / alternative-name lookup
   ↓
3. Word/less-strict candidate search
   ↓
4. Candidate CID collection
   ↓
5. Retrieve structure for each candidate
   ↓
6. RDKit validation
   ↓
7. Compare identifiers / requested name metadata
   ↓
Single confident candidate ──> continue
Multiple candidates ────────> ask user to select
No candidate ────────────────> clear resolution error
"""
)
doc.add_paragraph(
    "The application should never silently choose an arbitrary first result when a chemical name is ambiguous. "
    "A future resolver response should expose candidate CIDs, names and key descriptors when user selection is required."
)

# API
doc.add_heading("8. API Layer", level=1)
add_table(
    ["Method", "Endpoint", "Purpose"],
    [
        ["GET", "/", "Basic service information."],
        ["GET", "/health", "Health/status check."],
        ["POST", "/api/v1/structure/smiles", "Parse SMILES and return identifiers/descriptors."],
        ["POST", "/api/v1/structure/smiles/generate", "Generate complete structure package from SMILES."],
        ["POST", "/api/v1/structure/iupac", "Resolve IUPAC/chemical name and return identifiers/descriptors."],
        ["POST", "/api/v1/structure/iupac/generate", "Resolve name and generate complete structure package."],
        ["GET", "/api/v1/structure/{job_id}", "Retrieve stored generation metadata."],
        ["GET", "/api/v1/structure/{job_id}/download/{file_format}", "Download a generated file."],
        ["GET", "/api/v1/structure/{job_id}/json", "Download metadata JSON."],
    ],
    [0.8, 3.0, 3.7]
)

doc.add_heading("8.1 API Processing Logic", level=2)
add_code(
"""POST /api/v1/structure/iupac/generate
        ↓
Validate request body
        ↓
IUPACResolver.resolve(name)
        ↓
MoleculeEngine.from_smiles(resolved_smiles)
        ↓
generate_complete_structure()
        ├── get_identifiers()
        ├── get_properties()
        ├── render PNG
        ├── render SVG
        ├── generate 3D
        ├── export SDF
        ├── export MOL
        ├── export PDB
        ├── export XYZ
        └── write metadata.json
        ↓
JSON response with job_id + outputs
"""
)

# Frontend
doc.add_heading("9. Frontend Architecture", level=1)
doc.add_paragraph(
    "The frontend is implemented using React and Vite. The original scientific workstation-style UI is "
    "retained because it is the preferred interface design. The frontend communicates with FastAPI through "
    "a small API service layer rather than embedding backend logic in components."
)
add_table(
    ["File", "Role"],
    [
        ["frontend/src/App.jsx", "Main application interface, input modes, generation state and result presentation."],
        ["frontend/src/App.css", "Scientific workstation visual design and layout."],
        ["frontend/src/index.css", "Global/base styles."],
        ["frontend/src/main.jsx", "React application entry point."],
        ["frontend/src/services/api.js", "HTTP calls to FastAPI and generated-file URL handling."],
    ],
    [2.5, 5.0]
)

doc.add_heading("9.1 Frontend API Flow", level=2)
add_code(
"""React input
   ↓
api.js
   ↓
fetch()
   ↓
FastAPI endpoint
   ↓
JSON response
   ↓
React state
   ↓
Structure + descriptors + download links
"""
)

# File outputs
doc.add_heading("10. Output and Export System", level=1)
add_table(
    ["Output", "Current status", "Source"],
    [
        ["2D PNG", "Implemented", "RDKit 2D renderer"],
        ["2D SVG", "Implemented", "RDKit MolDraw2DSVG"],
        ["3D SDF", "Implemented", "RDKit SDWriter"],
        ["3D MOL", "Implemented", "RDKit MolToMolFile"],
        ["3D PDB", "Implemented", "RDKit MolToPDBFile"],
        ["3D XYZ", "Implemented", "Custom XYZ writer from conformer coordinates"],
        ["Metadata JSON", "Implemented", "FastAPI job metadata"],
        ["CML", "Planned", "Future exporter"],
        ["PDF report", "Planned", "Future report generator"],
        ["TIFF", "Planned", "Future image export"],
        ["WEBP", "Planned", "Future image export"],
        ["CSV", "Planned", "Future descriptor/export layer"],
        ["ZIP bundle", "Planned", "Future package export"],
    ],
    [2.0, 1.4, 4.1]
)

# Data model
doc.add_heading("11. Job and Data Model", level=1)
add_code(
"""exports/
└── {job_id}/
    ├── structure_2d.png
    ├── structure_2d.svg
    ├── structure_3d.sdf
    ├── structure_3d.mol
    ├── structure_3d.pdb
    ├── structure_3d.xyz
    └── metadata.json
"""
)
doc.add_paragraph(
    "The UUID job identifier creates a clean separation between independent generation requests. "
    "metadata.json records the input type/value, identifiers, molecular properties, 2D paths, 3D paths, "
    "optimization method, energy, atom count and conformer count."
)

# Validation
doc.add_heading("12. Chemical Validation Strategy", level=1)
add_bullets([
    "Input emptiness validation is performed before chemical processing.",
    "SMILES parsing is performed by RDKit.",
    "RDKit sanitization is used to reject chemically invalid molecular graphs.",
    "Identifiers and descriptors are calculated only after successful parsing.",
    "IUPAC-resolved structures are also passed through RDKit rather than trusted blindly.",
    "Future image-derived structures must pass the same RDKit validation gate before being treated as valid outputs.",
    "Future production validation should include valence checks, disconnected-component handling, stereochemistry checks and explicit ambiguity reporting."
])

# Example validation
doc.add_heading("13. Verified Test Case — Aspirin", level=1)
doc.add_paragraph("Input SMILES:")
add_code("CC(=O)OC1=CC=CC=C1C(=O)O")
add_table(
    ["Descriptor", "Verified value"],
    [
        ["Formula", "C9H8O4"],
        ["Molecular weight", "180.159"],
        ["Exact mass", "180.042259"],
        ["LogP", "1.3101"],
        ["TPSA", "63.6"],
        ["HBD", "1"],
        ["HBA", "3"],
        ["Rotatable bonds", "2"],
        ["Heavy atoms", "13"],
        ["Rings", "1"],
        ["Aromatic rings", "1"],
        ["Formal charge", "0"],
        ["Canonical SMILES", "CC(=O)Oc1ccccc1C(=O)O"],
        ["InChIKey", "BSYNRYMUTXBXSQ-UHFFFAOYSA-N"],
    ],
    [2.8, 4.7]
)

doc.add_heading("14. Verified Test Case — Complex IUPAC Input", level=1)
doc.add_paragraph("Tested chemical name:")
add_code(
"[1-[6-[[5-[4-(dimethylamino)piperidin-1-yl]pyridin-2-yl]amino]-"
"2-(4-fluoroanilino)pyrido[3,4-d]pyrimidin-4-yl]piperidin-4-yl]methanol"
)
add_table(
    ["Descriptor", "Observed value"],
    [
        ["Formula", "C31H38FN9O"],
        ["Heavy atoms", "42"],
        ["Rings", "6"],
        ["Aromatic rings", "4"],
        ["Exact mass", "571.318335"],
        ["LogP", "4.7852"],
        ["TPSA", "105.57"],
        ["HBD", "3"],
        ["HBA", "10"],
        ["Rotatable bonds", "8"],
        ["Formal charge", "0"],
    ],
    [2.8, 4.7]
)
doc.add_paragraph(
    "This test demonstrates that a long, structurally complex chemical name can be resolved and passed "
    "through the same molecular analysis and 2D rendering pipeline. The remaining engineering requirement is "
    "to make the resolver robust to naming variants and ambiguous names rather than relying on a single exact lookup."
)

# Engine design
doc.add_heading("15. Chemistry Engine Design Principles", level=1)
add_bullets([
    "One molecular graph should be the central internal representation.",
    "Input-specific modules should only be responsible for obtaining that molecular graph.",
    "RDKit should remain the deterministic chemistry authority for parsing, sanitization, descriptors and structure conversion.",
    "Rendering should consume validated molecules and never raw user strings.",
    "3D generation should be explicitly described as computational conformer generation.",
    "Exports should be generated from the same validated structure to prevent representation drift.",
    "Every generated job should have a stable job_id and machine-readable metadata."
])

# ML
doc.add_heading("16. Image-to-Structure ML Module — Planned Architecture", level=1)
doc.add_paragraph(
    "The ML module is intentionally isolated from the deterministic chemistry core. Its task is to infer "
    "a molecular representation from a 2D chemical structure image."
)
add_code(
"""Image
  ↓
Preprocessing
  ├── crop / deskew
  ├── denoise
  ├── resize / normalize
  └── background handling
  ↓
Vision model
  ↓
Predicted SMILES / molecular graph
  ↓
RDKit parse + sanitization
  ↓
Chemical validation
  ↓
Confidence / uncertainty
  ↓
User-facing structure + exports
"""
)
doc.add_paragraph(
    "A production image pipeline should not display an ML prediction as chemically correct merely because the "
    "model produced a string. The predicted structure must pass the same chemistry validation layer and should "
    "carry confidence/provenance information."
)

# Explainability
doc.add_heading("17. Validation, Confidence and Explainability — Planned", level=1)
add_bullets([
    "Input validity state.",
    "Resolver source and source identifier.",
    "Canonical structure identifier such as InChIKey.",
    "Chemical sanitization result.",
    "Image-model confidence for image inputs.",
    "Ambiguity flags for names resolving to multiple candidates.",
    "Provenance showing which module produced each representation.",
    "Warnings when 3D geometry is computational rather than experimentally determined."
])

# Future roadmap
doc.add_heading("18. Development Roadmap", level=1)
add_table(
    ["Phase", "Work", "Expected outcome"],
    [
        ["Phase 1", "Core chemistry", "RDKit parsing, identifiers and descriptors."],
        ["Phase 2", "2D rendering", "PNG/SVG structure images."],
        ["Phase 3", "3D engine", "ETKDGv3 + MMFF/UFF + 3D exports."],
        ["Phase 4", "FastAPI", "API endpoints and job-based outputs."],
        ["Phase 5", "IUPAC", "Name-to-structure resolution."],
        ["Phase 6", "IUPAC hardening", "Synonyms, variants, ambiguity handling and provenance."],
        ["Phase 7", "Image ML", "Image-to-structure inference."],
        ["Phase 8", "Interactive 3D", "Browser-based molecular viewer."],
        ["Phase 9", "Extended exports", "CML, PDF, TIFF, WEBP, CSV, ZIP."],
        ["Phase 10", "Validation/explainability", "Confidence, provenance and warnings."],
        ["Phase 11", "History/library", "Persistent structure records and reusable jobs."],
        ["Phase 12", "Production", "Security, rate limiting, queues, caching, monitoring and deployment."],
    ],
    [1.0, 3.0, 3.7]
)

# Production hardening
doc.add_heading("19. Production / Commercial-Grade Requirements", level=1)
add_bullets([
    "Move long-running 3D and ML jobs to a background task queue.",
    "Add structured application logging and request/job tracing.",
    "Add API rate limiting.",
    "Add request size and image upload limits.",
    "Validate file types rather than trusting file extensions.",
    "Add database-backed job records for persistent history.",
    "Cache repeated name-resolution requests where appropriate.",
    "Add deterministic version metadata: application version, RDKit version and model version.",
    "Add unit tests for each chemistry module and API endpoint.",
    "Add integration tests covering complete input-to-export workflows.",
    "Add security controls for uploaded images and generated files.",
    "Separate temporary files from persistent user files.",
    "Add Docker-based reproducible deployment.",
    "Add CI tests before release."
])

# Testing matrix
doc.add_heading("20. Recommended Testing Matrix", level=1)
add_table(
    ["Test category", "Examples"],
    [
        ["Valid SMILES", "Aspirin, ethanol, caffeine, complex heterocycles."],
        ["Invalid SMILES", "Empty string, malformed ring closure, invalid valence."],
        ["Valid names", "Common names, systematic names, long names."],
        ["Name variants", "Alternative nomenclature, synonyms, punctuation variants."],
        ["Ambiguous names", "Names returning multiple candidate records."],
        ["Charged molecules", "Ions and salts."],
        ["Stereochemistry", "Isomeric SMILES and chiral structures."],
        ["3D failures", "Molecules where embedding/optimization requires fallback."],
        ["Export integrity", "Open generated SDF/MOL/PDB/XYZ and verify atom/bond counts."],
        ["API errors", "404, 400, resolver timeout, malformed request."],
        ["Frontend integration", "Input → API → results → download."],
        ["Image ML", "Correct, noisy, rotated and partially cropped structure images."],
    ],
    [2.0, 5.7]
)

# Current limitations
doc.add_heading("21. Current Limitations", level=1)
add_bullets([
    "IUPAC resolution currently depends on external PubChem availability and response behavior.",
    "The resolver is not yet a multi-source chemical name reconciliation system.",
    "Ambiguous names are not yet exposed as a candidate-selection workflow.",
    "Image-to-structure ML is not yet integrated.",
    "The frontend currently focuses on the existing workstation UI rather than a full interactive 3D molecular viewer.",
    "Only the currently implemented export formats are available through the API.",
    "There is no persistent database-backed user/job history yet.",
    "The current 3D generation is conformer generation and force-field optimization, not experimental 3D structure determination."
])

# Final architecture snapshot
doc.add_heading("22. Final Target Architecture", level=1)
add_code(
"""                    ┌─────────────────────────────┐
                    │      React Frontend         │
                    │ SMILES / IUPAC / Image / 3D │
                    └──────────────┬──────────────┘
                                   │
                              REST / JSON
                                   │
                    ┌──────────────▼──────────────┐
                    │          FastAPI             │
                    │ Validation / Jobs / Exports  │
                    └──────────────┬──────────────┘
                                   │
                         ┌─────────▼─────────┐
                         │  MoleculeEngine   │
                         │     RDKit         │
                         └─────────┬─────────┘
                                   │
              ┌────────────────────┼────────────────────┐
              │                    │                    │
       Name Resolver          2D Renderer          3D Engine
       PubChem/others          PNG/SVG           ETKDG/MMFF/UFF
              │                    │                    │
              └────────────────────┼────────────────────┘
                                   │
                         ┌─────────▼─────────┐
                         │ Validation Layer  │
                         │ IDs / descriptors │
                         │ confidence / QA   │
                         └─────────┬─────────┘
                                   │
                  ┌────────────────┼─────────────────┐
                  │                │                 │
               Exports          Job Store        ML Module
            SDF/MOL/PDB/      metadata/history    Image→Mol
             XYZ/PDF/etc.
"""
)

# Conclusion
doc.add_heading("23. Conclusion", level=1)
doc.add_paragraph(
    "MolecularStructureAI has progressed from a basic molecular parsing concept into a modular chemistry "
    "application with a working RDKit engine, 2D/3D generation, FastAPI backend, IUPAC resolution and React "
    "frontend. The key architectural decision is the use of a single validated molecular graph as the common "
    "internal representation. This makes future input methods—especially image-to-structure ML—compatible with "
    "the same downstream analysis and export pipeline."
)
doc.add_paragraph(
    "The immediate engineering priority is not a UI redesign. It is backend robustness: strengthen the chemical "
    "name resolver, introduce ambiguity handling and provenance, then implement image-to-structure ML and the "
    "remaining export/viewer capabilities. This preserves the existing interface while progressively moving the "
    "system toward a research-grade and eventually commercial-grade molecular informatics platform."
)

# Footer
for section in doc.sections:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("MolecularStructureAI — Detailed Technical Project Report").font.size = Pt(8)

doc.save(out)
print(f"Created: {out}")
